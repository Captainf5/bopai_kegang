"""Cloud rendering only: no model calls and no execution of submitted Markdown."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import layout_identity


ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "outputs/catalog/博AI增效-2D-10倍职场办公_6X畅销版.md"
TEMPLATE = ROOT / "delivery-pack/课纲模板.docx"
RENDERER = Path(__file__).with_name("课纲排版toWord带图版.py")
ACK = "我确认课纲不含客户机密或个人敏感信息，并同意公开课纲与生成文件。"
REQUIRED_SECTIONS = [
    "主讲人介绍",
    "本课程说明",
    "课程大纲",
    "课程产出",
    "课件展示",
    "课程现场",
    "课后准备",
    "培训准备",
]
MODULE_FIELDS = ["关键词：", "讲解要点：", "演示：", "学员练习：", "产出物："]
OPC_TITLE = "# 博AI增效-OPC实战工作坊"
COURSE_TITLE_RE = re.compile(
    r"^# 【博AI增效-(?P<duration>(?:[1-9]\d*D|(?:[1-9]\d*(?:\.\d+)?|0\.\d+)H))】(?P<course_name>.+)$"
)
PRESERVED_TEMPLATE_PARTS = [
    "word/styles.xml",
    "word/stylesWithEffects.xml",
    "word/numbering.xml",
    "word/theme/theme1.xml",
]


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def display_path(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return path.name


def issue_markdown(event):
    owner = event["repository"]["owner"]["login"]
    issue = event["issue"]
    if event["sender"]["login"] != owner or issue["user"]["login"] != owner:
        raise ValueError("仅仓库所有者可以使用公开排版入口")
    if not issue["title"].startswith("[课纲排版]"):
        raise ValueError("不是课纲排版请求")
    body = issue.get("body") or ""
    if "- [x] " + ACK not in body.lower():
        raise ValueError("未确认公开范围")
    marker = "### 课纲 Markdown\n\n"
    if marker not in body or "\n### 发布确认" not in body:
        raise ValueError("请使用仓库课纲排版表单")
    content = body.split(marker, 1)[1].rsplit("\n### 发布确认", 1)[0].strip()
    fence = chr(96) * 3
    if content.startswith(fence) and content.endswith(fence):
        content = content.split("\n", 1)[1].rsplit(fence, 1)[0].strip()
    if not content.startswith("# ") or len(content) < 300 or len(content) > 60000:
        raise ValueError("请提交完整课纲Markdown（300—60000字符），而非一句需求")
    return content, int(issue["number"])


def validate_course_markdown(content: str) -> dict:
    """Apply the same minimum course contract before any layout is generated."""
    first_line = content.splitlines()[0].strip() if content.splitlines() else ""
    if first_line == OPC_TITLE:
        title_duration = "OPC"
    else:
        title_match = COURSE_TITLE_RE.fullmatch(first_line)
        if not title_match or not title_match.group("course_name").strip():
            raise ValueError(
                "课纲标题必须使用 # 【博AI增效-{时长}】{课程名}；"
                "小时课按实际时长标记，OPC课程仅允许精确标题 # 博AI增效-OPC实战工作坊"
            )
        title_duration = title_match.group("duration")

    headings = []
    for section in REQUIRED_SECTIONS:
        match = re.search(rf"^##\s+[^\n]*{re.escape(section)}\s*$", content, re.M)
        if not match:
            raise ValueError(f"课纲缺少必要章节：{section}")
        headings.append(match.start())
    if headings != sorted(headings):
        raise ValueError("课纲章节顺序不符合统一结构")

    module_matches = list(
        re.finditer(r"^\*\*(?:模块\s*\d+|加餐)：[^\n]+\*\*\s*$", content, re.M)
    )
    if not module_matches:
        raise ValueError("课程大纲中没有完整教学模块")
    for index, match in enumerate(module_matches):
        end = module_matches[index + 1].start() if index + 1 < len(module_matches) else len(content)
        block = content[match.end():end]
        positions = []
        for field in MODULE_FIELDS:
            position = block.find(field)
            if position < 0:
                raise ValueError(f"{match.group(0)} 缺少字段：{field}")
            positions.append(position)
        if positions != sorted(positions):
            raise ValueError(f"{match.group(0)} 的模块字段顺序错误")

    rows = re.findall(r"^\|\s*(\d{2}:\d{2})—(\d{2}:\d{2})\s*\|", content, re.M)
    slots = re.findall(
        r"^\*\*(?:模块\s*\d+|加餐)：[^\n]*｜(\d{2}:\d{2})—(\d{2}:\d{2})\*\*\s*$",
        content,
        re.M,
    )
    if not rows or rows != slots:
        raise ValueError("课程时间表必须与教学模块时段逐项一致")

    return {
        "title_duration": title_duration,
        "required_sections": len(REQUIRED_SECTIONS),
        "modules": len(module_matches),
        "schedule_rows": len(rows),
    }


def resolve_soffice(soffice: str) -> str:
    binary = shutil.which(soffice) if not Path(soffice).is_absolute() else soffice
    if not binary:
        raise RuntimeError("缺少云端LibreOffice；请使用GitHub Actions排版入口")
    return str(binary)


def libreoffice_version(binary: str) -> str:
    result = subprocess.run(
        [binary, "--version"],
        capture_output=True,
        text=True,
        timeout=30,
    )
    return (result.stdout or result.stderr).strip().splitlines()[0][:200]


def convert_pdf(docx: Path, pdf: Path, soffice: str) -> str:
    binary = resolve_soffice(soffice)
    with tempfile.TemporaryDirectory(prefix="bopai-render-") as tmp_name:
        tmp = Path(tmp_name)
        result = subprocess.run(
            [
                binary,
                "-env:UserInstallation=" + (tmp / "profile").as_uri(),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(tmp),
                str(docx.resolve()),
            ],
            capture_output=True,
            text=True,
            timeout=180,
        )
        generated = tmp / (docx.stem + ".pdf")
        if result.returncode != 0 or not generated.exists() or generated.stat().st_size < 1000:
            raise RuntimeError("PDF转换失败: " + result.stderr[-500:])
        shutil.copyfile(generated, pdf)
    return binary


def audit_docx(docx_path: Path) -> dict:
    document = Document(docx_path)
    if not document.sections:
        raise ValueError("Word版式检查失败：缺少页面设置")
    for section in document.sections:
        if round(section.page_width.pt) != 612 or round(section.page_height.pt) != 792:
            raise ValueError("Word版式检查失败：页面不是统一Letter尺寸")

    with zipfile.ZipFile(docx_path) as actual, zipfile.ZipFile(TEMPLATE) as expected:
        for part in PRESERVED_TEMPLATE_PARTS:
            if part in expected.namelist() and actual.read(part) != expected.read(part):
                raise ValueError(f"Word版式检查失败：模板部件发生漂移 {part}")

    image_rids = {
        rid for rid, relationship in document.part.rels.items() if relationship.reltype == RT.IMAGE
    }
    used_rids = {
        value
        for element in document._element.iter()
        for attribute, value in element.attrib.items()
        if attribute in {qn("r:embed"), qn("r:link"), qn("r:id")}
    }
    orphan_images = sorted(image_rids - used_rids)
    if orphan_images:
        raise ValueError("Word版式检查失败：存在未引用图片关系")

    return {
        "sections": len(document.sections),
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "inline_images": len(document.inline_shapes),
        "orphan_images": orphan_images,
        "page_size_points": [612, 792],
        "template_parts_preserved": True,
    }


def audit_pdf(pdf_path: Path) -> dict:
    import pypdfium2

    document = pypdfium2.PdfDocument(pdf_path)
    page_text_lengths = []
    for page in document:
        text = page.get_textpage().get_text_range()
        page_text_lengths.append(len(re.sub(r"\s+", "", text)))
    blank_pages = [index + 1 for index, length in enumerate(page_text_lengths) if length == 0]
    if blank_pages:
        raise ValueError(f"PDF版式检查失败：存在纯空白页 {blank_pages}")
    if not page_text_lengths or sum(page_text_lengths) == 0:
        raise ValueError("PDF正文检查失败")
    return {
        "pages": len(page_text_lengths),
        "page_text_lengths": page_text_lengths,
        "blank_pages": blank_pages,
        "low_text_pages_for_review": [
            index + 1 for index, length in enumerate(page_text_lengths) if length < 40
        ],
    }


def load_layout_module():
    spec = importlib.util.spec_from_file_location("layout", RENDERER)
    layout = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(layout)
    return layout


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--event", type=Path)
    parser.add_argument("--soffice", default="soffice")
    parser.add_argument("--input", type=Path)
    args = parser.parse_args()

    if args.event:
        content, number = issue_markdown(
            json.loads(args.event.read_text(encoding="utf-8"))
        )
        dest = ROOT / "outputs/requests" / ("issue-" + str(number))
        dest.mkdir(parents=True, exist_ok=True)
        md = dest / "course.md"
        md.write_text(content + "\n", encoding="utf-8")
    else:
        md = (args.input or SAMPLE).resolve()
        content = md.read_text(encoding="utf-8-sig")

    markdown_checks = validate_course_markdown(content)
    docx = md.with_name(md.stem + "_带图.docx")
    pdf = docx.with_suffix(".pdf")
    load_layout_module().convert_md_to_docx(md, docx)
    docx_checks = audit_docx(docx)
    binary = convert_pdf(docx, pdf, args.soffice)
    pdf_checks = audit_pdf(pdf)

    assets = layout_identity.layout_asset_hashes()
    manifest = {
        "status": "passed",
        "skill_version": (ROOT / "VERSION").read_text(encoding="utf-8").strip(),
        "github_sha": os.environ.get("GITHUB_SHA", ""),
        "files": [display_path(path) for path in [md, docx, pdf]],
        "source_sha256": sha256_file(md),
        "docx_sha256": sha256_file(docx),
        "pdf_sha256": sha256_file(pdf),
        "layout_basis_sha256": layout_identity.layout_basis_sha256(),
        "layout_assets": assets,
        "libreoffice_version": libreoffice_version(binary),
        "markdown_checks": markdown_checks,
        "docx_checks": docx_checks,
        "pdf_checks": pdf_checks,
        "pages": pdf_checks["pages"],
    }
    Path("render-result.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(manifest, ensure_ascii=False))


if __name__ == "__main__":
    main()
