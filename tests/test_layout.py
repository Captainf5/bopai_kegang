import importlib.util, re, tempfile, unittest, zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT = Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location("layout",ROOT/"delivery-pack/scripts/课纲排版toWord带图版.py")
layout=importlib.util.module_from_spec(spec); spec.loader.exec_module(layout)
class LayoutTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory()
        cls.out=Path(cls.tmp.name)/"test.docx"
        layout.convert_md_to_docx(ROOT/"outputs/catalog/博AI增效-2D-10倍职场办公_6X畅销版.md",cls.out)
        cls.doc=Document(cls.out)
    @classmethod
    def tearDownClass(cls): cls.tmp.cleanup()
    def test_day_two_new_page(self):
        day=next(p for p in self.doc.paragraphs if p.text.startswith("第二天"))
        self.assertTrue(day.paragraph_format.page_break_before)
    def test_only_new_training_days_force_page_breaks(self):
        ordinary_sections=['三、课程大纲','四、课程产出','五、课件展示','六、课程现场','七、课后准备','八、培训准备']
        for title in ordinary_sections:
            paragraph=next(p for p in self.doc.paragraphs if p.text==title)
            self.assertFalse(bool(paragraph.paragraph_format.page_break_before),title)
        day=next(p for p in self.doc.paragraphs if p.text.startswith("第二天"))
        self.assertTrue(day.paragraph_format.page_break_before)
    def test_body_paragraphs_may_flow_across_pages(self):
        body=next(p for p in self.doc.paragraphs if p.text.startswith("讲解要点："))
        self.assertFalse(bool(body.paragraph_format.keep_together))
    def test_module_headings_not_orphaned(self):
        modules=[p for p in self.doc.paragraphs if re.match(r"模块\d+：",p.text)]
        self.assertEqual(len(modules),8)
        self.assertTrue(all(p.paragraph_format.keep_with_next for p in modules))
    def test_heading_outline_semantics(self):
        def outline(paragraph):
            values=paragraph._p.xpath('./w:pPr/w:outlineLvl/@w:val')
            return values[0] if values else None
        self.assertEqual(outline(self.doc.paragraphs[0]),'0')
        chapters=[p for p in self.doc.paragraphs if re.match(r'^[一二三四五六七八]、',p.text)]
        self.assertEqual(len(chapters),8)
        self.assertTrue(all(outline(p)=='1' for p in chapters))
        days=[p for p in self.doc.paragraphs if p.text.startswith(('第一天','第二天'))]
        self.assertEqual(len(days),2)
        self.assertTrue(all(outline(p)=='2' for p in days))
    def test_long_title_uses_compact_font_only_when_needed(self):
        doc=Document()
        layout.add_h1(doc,'博AI增效-2D-通讯Token运营AI实战工作坊_6X垂直版')
        self.assertEqual(doc.paragraphs[-1].runs[0].font.size.pt,16)
        layout.add_h1(doc,'博AI增效-1D-10倍职场办公_6X畅销版')
        self.assertEqual(doc.paragraphs[-1].runs[0].font.size.pt,18)
    def test_images_only_in_correct_sections(self):
        self.assertEqual(len(self.doc.inline_shapes),6)
        image_rels={rid for rid,rel in self.doc.part.rels.items() if rel.reltype==RT.IMAGE}
        used_rids={
            element.get(qn('r:embed'))
            for element in self.doc._element.iter()
            if element.get(qn('r:embed'))
        }
        self.assertEqual(image_rels,used_rids)
        with zipfile.ZipFile(self.out) as archive:
            self.assertNotIn('word/media/image7.png',archive.namelist())
    def test_geometry_and_template_preservation(self):
        self.assertEqual(len(self.doc.sections),1)
        self.assertEqual(self.doc.sections[0].page_width.pt,612)
        with zipfile.ZipFile(self.out) as actual,zipfile.ZipFile(ROOT/"delivery-pack/课纲模板.docx") as ref:
            for p in ["word/styles.xml","word/numbering.xml","word/theme/theme1.xml"]:
                self.assertEqual(actual.read(p),ref.read(p))
    def test_daily_timetables_and_breaks(self):
        schedules=[t for t in self.doc.tables if t.cell(0,0).text=="时间"]
        self.assertEqual(len(schedules),2)
        for table in schedules:
            last=None
            for row in table.rows[1:]:
                a,b=re.fullmatch(r"(\d{2}:\d{2})—(\d{2}:\d{2})",row.cells[0].text).groups()
                minutes=lambda s:int(s[:2])*60+int(s[3:])
                start,end=minutes(a),minutes(b)
                self.assertGreater(end,start)
                if last is not None:
                    self.assertIn(start-last,[15,120])
                last=end
            self.assertEqual(last,17*60+30)
    def test_reference_count(self):
        self.assertEqual(len(list((ROOT/"references").glob("[0-9][0-9]-*.md"))),10)
if __name__=="__main__": unittest.main()
