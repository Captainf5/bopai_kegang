import importlib.util, re, tempfile, unittest
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('variant_layout', ROOT/'delivery-pack/scripts/课纲排版toWord带图版.py')
layout = importlib.util.module_from_spec(spec)
spec.loader.exec_module(layout)

class VariantTests(unittest.TestCase):
    def test_keywords_and_timetables(self):
        for days, count in [(1, 6), (2, 9)]:
            with self.subTest(days=days), tempfile.TemporaryDirectory() as tmp:
                md = ROOT/f'outputs/WorkBuddy职场办公实战工作坊-{days}D.md'
                source = md.read_text(encoding='utf-8')
                out = Path(tmp)/'course.docx'
                layout.convert_md_to_docx(md, out)
                doc = Document(out)
                keywords = [p for p in doc.paragraphs if p.text.startswith('关键词：')]
                self.assertEqual(len(keywords), count)
                self.assertTrue(all(p.paragraph_format.keep_with_next for p in keywords))
                self.assertEqual(source.count('- 讲解要点：'), count)
                modules = re.split(r'\*\*模块\d+：', source)[1:]
                self.assertIn('豆包', modules[0]); self.assertIn('DeepSeek', modules[0])
                for term in ['AGENTS.md', '连接器的使用', '专家团的创建', '自动化任务', '岗位AI工作台', 'Reconcile', 'OC框架']:
                    self.assertIn(term, source)
                tables = [t for t in doc.tables if t.cell(0,0).text == '时间']
                self.assertEqual(len(tables), days)
                for table in tables:
                    previous = None
                    for row in table.rows[1:]:
                        a,b = re.fullmatch(r'(\d{2}:\d{2})—(\d{2}:\d{2})',row.cells[0].text).groups()
                        minutes = lambda s: int(s[:2])*60+int(s[3:])
                        start,end = minutes(a),minutes(b)
                        self.assertGreater(end,start)
                        if previous is None: self.assertEqual(start,540)
                        else: self.assertIn(start-previous,[15,120])
                        previous=end
                    self.assertEqual(previous,1050)
